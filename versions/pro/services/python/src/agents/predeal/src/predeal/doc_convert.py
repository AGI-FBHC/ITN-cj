"""Office 文档 → Markdown。

只做“够用”的结构抽取：标题层级 + 段落 + 表格 / 幻灯片要点，
让下游切片与目录生成能识别章节。
DOCX 用标准库解析（zip + XML），PPTX 复用已安装的 python-pptx。
"""

from __future__ import annotations

import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

SUPPORTED_OFFICE_SUFFIXES = {".docx", ".pptx"}


def convert_to_markdown(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return _docx_to_markdown(path)
    if suffix == ".pptx":
        return _pptx_to_markdown(path)
    raise ValueError(f"Unsupported office file type: {suffix}")


def _heading_level(style_id: str, outline_level: str | None) -> int | None:
    """返回 1-6 的标题级别；不是标题返回 None。"""
    if outline_level is not None:
        try:
            level = int(outline_level) + 1
            if 1 <= level <= 6:
                return level
        except ValueError:
            pass
    lowered = (style_id or "").lower()
    if "heading" in lowered:
        digits = re.findall(r"\d+", lowered)
        if digits:
            return min(max(int(digits[0]), 1), 6)
        return 1
    if lowered in {"title"}:
        return 1
    return None


def _docx_to_markdown(path: Path) -> str:
    with zipfile.ZipFile(path) as archive:
        if "word/document.xml" not in archive.namelist():
            raise ValueError("DOCX 缺少 word/document.xml")
        root = ET.fromstring(archive.read("word/document.xml"))

    body = root.find(f"{_W}body")
    if body is None:
        return ""

    lines: list[str] = []
    for node in list(body):
        tag = node.tag
        if tag == f"{_W}p":
            text = "".join(t.text or "" for t in node.iter(f"{_W}t")).strip()
            if not text:
                continue
            style_id = ""
            outline_level: str | None = None
            style_node = node.find(f"{_W}pPr/{_W}pStyle")
            if style_node is not None:
                style_id = style_node.get(f"{_W}val", "") or ""
            outline_node = node.find(f"{_W}pPr/{_W}outlineLvl")
            if outline_node is not None:
                outline_level = outline_node.get(f"{_W}val")

            level = _heading_level(style_id, outline_level)
            if level is not None:
                lines.append("#" * level + " " + text)
            elif "list" in (style_id or "").lower():
                lines.append("- " + text)
            else:
                lines.append(text)
            lines.append("")
        elif tag == f"{_W}tbl":
            for row in node.findall(f"{_W}tr"):
                cells: list[str] = []
                for cell in row.findall(f"{_W}tc"):
                    cell_text = " ".join(
                        "".join(t.text or "" for t in para.iter(f"{_W}t")).strip()
                        for para in cell.findall(f"{_W}p")
                    ).strip()
                    cells.append(cell_text)
                if any(cells):
                    lines.append("| " + " | ".join(cells) + " |")
            lines.append("")

    return "\n".join(lines).strip() + "\n" if lines else ""


def _pptx_to_markdown(path: Path) -> str:
    from pptx import Presentation

    prs = Presentation(str(path))
    lines: list[str] = []
    for index, slide in enumerate(prs.slides, start=1):
        title = ""
        try:
            if slide.shapes.title is not None:
                title = (slide.shapes.title.text or "").strip()
        except Exception:
            title = ""
        lines.append("## " + (title or f"第 {index} 部分"))
        lines.append("")
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            for paragraph in shape.text_frame.paragraphs:
                text = "".join(run.text or "" for run in paragraph.runs).strip()
                if not text:
                    text = (paragraph.text or "").strip()
                if not text or text == title:
                    continue
                lines.append("- " + text)
        lines.append("")
    return "\n".join(lines).strip() + "\n" if lines else ""
