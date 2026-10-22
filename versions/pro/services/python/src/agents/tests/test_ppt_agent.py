import asyncio
import json
import re
from pathlib import Path
from types import SimpleNamespace

import pytest
from pptx import Presentation

import app.presentations.agent as presentation_agent_module
from app.presentations.agent import PptAgent, _fallback_outline
from app.presentations.backgrounds import (
    COVER_TEXT_OVERRIDES,
    THEME_PALETTES,
    palette_for,
    select_backgrounds,
)
from app.presentations.models import PresentationRequest


class NoopStore:
    pass


def test_teaching_moments_are_anchored_and_old_slides_get_runtime_fallbacks():
    slides = [
        {"title": "封面", "kind": "title", "speakerNotes": "欢迎学习。"},
        {
            "title": "核心机制", "kind": "concept",
            "points": [
                {"heading": "条件", "detail": "先识别适用条件，再选择执行方法。"},
                {"heading": "校验", "detail": "最后使用结果校验约束是否成立。"},
            ],
            "speakerNotes": "先识别适用条件，再选择执行方法。完成操作后，还要使用结果校验约束是否成立。",
            "teachingMoments": [{
                "type": "quiz", "anchorText": "再选择执行方法", "title": "检查理解",
                "prompt": "应该先做什么？", "options": ["识别条件", "跳过检查", "直接套用", "忽略结果"],
                "answerIndex": 0, "explanation": "先判断适用条件。",
            }],
        },
        {"title": "总结", "kind": "conclusion", "speakerNotes": "回顾本节。"},
    ]

    PptAgent._ensure_teaching_moments(slides)

    moment = slides[1]["teachingMoments"][0]
    assert moment["momentId"] == "s2-m1"
    assert moment["anchorChar"] > 0
    assert 0.12 <= moment["anchorRatio"] <= 0.98
    assert slides[0]["teachingMoments"] == []


def test_teaching_moment_quiz_sanitizes_missing_stem_and_markdown_heading():
    slides = [
        {
            "title": "核心机制", "kind": "concept", "points": [],
            "speakerNotes": "先识别适用条件，再选择执行方法。完成操作后，还要使用结果校验约束是否成立。",
            "teachingMoments": [{
                "type": "quiz", "anchorText": "再选择执行方法", "title": "###", "prompt": "",
                "options": ["识别条件", "跳过检查", "直接套用", "忽略结果"],
                "answerIndex": 0, "explanation": "先判断适用条件。",
            }],
        },
        {"title": "总结", "kind": "conclusion", "speakerNotes": "回顾本节。"},
    ]

    PptAgent._ensure_teaching_moments(slides)

    moment = slides[0]["teachingMoments"][0]
    assert moment["title"] == "课堂检查点"
    assert moment["prompt"] == "请根据刚才讲到的内容，选择最准确的一项。"
    assert len(moment["options"]) == 4


class SizePlanningLlm:
    settings = SimpleNamespace(llm_timeout_seconds=60)

    async def complete_json(self, system, prompt, fallback, provider="", **kwargs):
        assert "[12, 20, 30, 40, 55, 70]" in prompt
        return {"slideCount": 55}


class TutorAgendaLlm:
    settings = SimpleNamespace(llm_timeout_seconds=60)

    async def complete_json(self, system, prompt, fallback, provider="", **kwargs):
        return {
            "slideCount": 30,
            "chapters": [
                {
                    "title": "第一章：命题逻辑",
                    "objective": "建立形式化推理基础",
                    "keyTopics": ["命题与联结词", "范式"],
                },
                {
                    "title": "集合与关系",
                    "objective": "理解离散结构的表达方法",
                    "keyTopics": ["集合运算", "等价关系"],
                },
                {
                    "title": "图论方法",
                    "objective": "使用图模型解决实际问题",
                    "keyTopics": ["路径与连通", "树"],
                },
            ],
        }


class RequirementEchoAgendaLlm:
    settings = SimpleNamespace(llm_timeout_seconds=60)

    async def complete_json(self, system, prompt, fallback, provider="", **kwargs):
        return {
            "slideCount": 12,
            "chapters": [
                {"title": "内容准确", "objective": "内容准确", "keyTopics": ["内容准确"]},
                {"title": "结构清晰", "objective": "结构清晰", "keyTopics": ["结构清晰"]},
                {"title": "自主学习", "objective": "自主学习", "keyTopics": ["自主学习"]},
            ],
        }


class ConcurrentBatchLlm:
    settings = SimpleNamespace(
        llm_timeout_seconds=60,
        presentation_content_min_concurrency=4,
        presentation_content_initial_concurrency=6,
        presentation_content_max_concurrency=8,
        presentation_content_retry_budget=2,
        presentation_content_batch_size=1,
    )

    def __init__(self):
        self.active = 0
        self.max_active = 0
        self.calls = 0

    async def complete_json(self, system, prompt, fallback, provider="", **kwargs):
        self.calls += 1
        self.active += 1
        self.max_active = max(self.max_active, self.active)
        await asyncio.sleep(.01)
        self.active -= 1
        seeds = json.loads(prompt.split("本批种子结构：", 1)[1])
        return {"slides": [
            {
                **seed,
                "title": seed.get("title", "页面"),
                "points": [
                    {"heading": "概念", "detail": f"{seed.get('title', '本页主题')}需要结合适用条件、执行方法和具体示例理解。"},
                    {"heading": "分析", "detail": "说明适用条件、执行步骤和常见问题。"},
                    {"heading": "实践", "detail": "给出课堂练习以及可验证的学习结果。"},
                ],
                "speakerNotes": "逐页讲解备注。",
            }
            for seed in seeds
        ]}


class SameModelRetryBatchLlm(ConcurrentBatchLlm):
    settings = SimpleNamespace(
        llm_timeout_seconds=60,
        presentation_content_min_concurrency=2,
        presentation_content_initial_concurrency=3,
        presentation_content_max_concurrency=4,
        presentation_content_retry_budget=2,
        presentation_content_batch_size=2,
        presentation_content_timeout_seconds=50,
    )

    async def complete_json(self, system, prompt, fallback, provider="", **kwargs):
        seeds = json.loads(prompt.split("本批种子结构：", 1)[1])
        if len(seeds) > 1:
            self.calls += 1
            return {"slides": []}
        return await super().complete_json(system, prompt, fallback, provider=provider, **kwargs)


class FailedPrimaryAndRetryLlm(SameModelRetryBatchLlm):
    async def complete_json(self, system, prompt, fallback, provider="", **kwargs):
        self.calls += 1
        return {"slides": []}


class RepairableInvalidBatchLlm(SameModelRetryBatchLlm):
    async def complete_json(self, system, prompt, fallback, provider="", **kwargs):
        self.calls += 1
        seeds = json.loads(prompt.split("本批种子结构：", 1)[1])
        if "修复同一份演示文稿" not in prompt:
            return {"slides": [
                {
                    "title": seed.get("title"),
                    "kind": seed.get("kind"),
                    "speakerNotes": "已有讲解，但正文结构损坏。",
                }
                for seed in seeds
            ]}
        return {"slides": [
            {
                **seed,
                "title": seed.get("title", "页面"),
                "points": [
                    {"heading": "定义", "detail": f"围绕{seed.get('title', '本页主题')}给出准确概念、适用条件和符号含义。"},
                    {"heading": "方法", "detail": "按照明确步骤完成推导，并检查每一步的成立条件。"},
                    {"heading": "例题", "detail": "通过具体计算验证方法，同时指出常见错误。"},
                ],
                "speakerNotes": "先说明定义，再结合计算步骤和例题完成这一页的讲解。",
            }
            for seed in seeds
        ]}


class RepeatedGenericBatchLlm(SameModelRetryBatchLlm):
    async def complete_json(self, system, prompt, fallback, provider="", **kwargs):
        self.calls += 1
        seeds = json.loads(prompt.split("本批种子结构：", 1)[1])
        return {"slides": [
            {
                "title": seed.get("title"),
                "kind": seed.get("kind"),
                "takeaway": "掌握核心知识并能够完成实际应用。",
                "points": [
                    {"heading": "定义", "detail": "给出准确概念、适用条件和符号含义。"},
                    {"heading": "方法", "detail": "按照明确步骤完成分析，并检查每一步是否正确。"},
                    {"heading": "例题", "detail": "通过具体练习验证方法，同时指出常见错误。"},
                ],
                "speakerNotes": "先说明定义，再结合方法和例题完成这一页的讲解。",
            }
            for seed in seeds
        ]}


class SchemaDriftBatchLlm(ConcurrentBatchLlm):
    settings = SameModelRetryBatchLlm.settings

    async def complete_json(self, system, prompt, fallback, provider="", **kwargs):
        seeds = json.loads(prompt.split("本批种子结构：", 1)[1])
        return {"slides": [
            {
                "title": seed.get("title"),
                "kind": "content",
                "points": [
                    {"title": "核心内容", "content": f"围绕{seed.get('title')}给出具体解释和案例，不应因字段名偏差被丢弃。"},
                    {"label": "应用方法", "description": "先判断条件，再选择方法，最后检查计算结果。"},
                    {"heading": "常见错误", "detail": "不要忽略定义域和边界条件，应在完成后进行验证。"},
                ],
            }
            for seed in seeds
        ]}


class FakeEnhancedTutor:
    def __init__(self, *args, **kwargs):
        pass

    async def stream(self, messages):
        assert "知识库或web_search工具" in messages[0]["content"]
        yield {"type": "tool_call_completed", "tool": "web_search", "success": True}
        yield {
            "type": "content",
            "text": json.dumps({
                "slideCount": 30,
                "chapters": [
                    {
                        "title": "第一章：命题逻辑",
                        "slides": [
                            {"title": "命题与联结词", "content": "建立形式化推理基础"},
                            {"title": "范式", "content": "掌握等值演算"},
                        ],
                    },
                    {
                        "title": "集合与关系",
                        "slides": [
                            {"title": "集合运算", "content": "理解离散结构的表达方法"},
                            {"title": "等价关系", "content": "识别关系性质"},
                        ],
                    },
                    {
                        "title": "图论方法",
                        "slides": [
                            {"title": "路径与连通", "content": "使用图模型解决实际问题"},
                            {"title": "树", "content": "掌握典型图结构"},
                        ],
                    },
                ],
            }, ensure_ascii=False),
        }


def make_request(**overrides):
    payload = {
        "requirement": "面向软件工程专业学生介绍 AI Agent 的架构、开发流程与实践案例",
        "title": "AI Agent 技术与实践",
        "audience": "students", "purpose": "courseware", "style": "technology",
        "backgroundStyle": "none",
        "language": "zh-CN", "slideCount": 8, "aspectRatio": "16:9",
        "detailLevel": "balanced", "includeSpeakerNotes": True, "modelProvider": "mimo",
    }
    payload.update(overrides)
    return PresentationRequest.model_validate(payload)


def test_content_batches_run_in_parallel_and_emit_metrics(tmp_path: Path):
    request = make_request(slideCount=12)
    llm = ConcurrentBatchLlm()
    agent = PptAgent(llm, NoopStore(), tmp_path)

    result = asyncio.run(agent._generate_in_batches(
        request, _fallback_outline(request), timeout=60,
    ))

    metrics = result["_generationMetrics"]
    assert len(result["slides"]) == 12
    assert metrics["batchSize"] == 1
    assert metrics["requestCount"] == 12
    assert metrics["retryCount"] == 0
    assert metrics["fallbackBatchCount"] == 0
    assert metrics["maxObservedConcurrency"] > 1
    assert llm.max_active > 1


def test_failed_batches_retry_individual_pages_with_the_same_model(tmp_path: Path):
    request = make_request(slideCount=6)
    llm = SameModelRetryBatchLlm()
    agent = PptAgent(llm, NoopStore(), tmp_path)

    result = asyncio.run(agent._generate_in_batches(
        request, _fallback_outline(request), timeout=60,
    ))

    metrics = result["_generationMetrics"]
    assert len(result["slides"]) == 6
    assert metrics["batchSize"] == 2
    assert metrics["retryCount"] == 6
    assert metrics["fallbackSlideCount"] == 0
    assert all(
        "围绕" not in point["detail"]
        for slide in result["slides"]
        for point in slide.get("points", [])
        if isinstance(point, dict)
    )


def test_schema_drift_is_repaired_without_retrying_or_switching_models(tmp_path: Path):
    request = make_request(
        requirement="面向大学生介绍时间管理方法与实践案例",
        title="时间管理",
        slideCount=6,
    )
    llm = SchemaDriftBatchLlm()
    result = asyncio.run(PptAgent(llm, NoopStore(), tmp_path)._generate_in_batches(
        request, _fallback_outline(request), timeout=60,
    ))

    metrics = result["_generationMetrics"]
    assert metrics["retryCount"] == 0
    assert metrics["fallbackSlideCount"] == 0
    assert all(slide.get("speakerNotes") for slide in result["slides"])
    assert all(slide.get("kind") != "content" for slide in result["slides"])


def test_generation_uses_local_fallback_when_mimo_retry_fails(tmp_path: Path):
    request = make_request(slideCount=6)
    agent = PptAgent(FailedPrimaryAndRetryLlm(), NoopStore(), tmp_path)

    result = asyncio.run(agent._generate_in_batches(
        request, _fallback_outline(request), timeout=60,
    ))

    metrics = result["_generationMetrics"]
    assert len(result["slides"]) == 6
    assert metrics["fallbackSlideCount"] > 0
    assert all(slide.get("speakerNotes") for slide in result["slides"])


def test_invalid_single_pages_are_repaired_by_mimo_instead_of_aborting_deck(tmp_path: Path):
    request = make_request(slideCount=6)
    llm = RepairableInvalidBatchLlm()

    result = asyncio.run(PptAgent(llm, NoopStore(), tmp_path)._generate_in_batches(
        request, _fallback_outline(request), timeout=60,
    ))

    metrics = result["_generationMetrics"]
    assert len(result["slides"]) == 6
    assert metrics["retryCount"] == 12
    assert metrics["fallbackSlideCount"] == 0
    assert all(len(slide["points"]) == 3 for slide in result["slides"])


def test_repeated_generic_recovery_content_does_not_fail_deck(tmp_path: Path):
    request = make_request(slideCount=6)

    result = asyncio.run(PptAgent(
        RepeatedGenericBatchLlm(), NoopStore(), tmp_path,
    )._generate_in_batches(
        request, _fallback_outline(request), timeout=60,
    ))
    assert len(result["slides"]) == 6


@pytest.mark.parametrize("retry_budget", [3, 4])
def test_configured_content_retry_budget_is_fully_used(tmp_path: Path, retry_budget: int):
    request = make_request(slideCount=6)
    llm = FailedPrimaryAndRetryLlm()
    settings = vars(SameModelRetryBatchLlm.settings).copy()
    settings["presentation_content_retry_budget"] = retry_budget
    llm.settings = SimpleNamespace(**settings)

    result = asyncio.run(PptAgent(llm, NoopStore(), tmp_path)._generate_in_batches(
        request, _fallback_outline(request), timeout=60,
    ))

    # 本地兜底保证整份课件仍可交付，不再因单页空返回中断整个任务。
    assert len(result["slides"]) == 6
    assert result["_generationMetrics"]["fallbackSlideCount"] == 6
    assert llm.calls > 0


def test_render_creates_valid_widescreen_presentation(tmp_path: Path):
    request = make_request()
    output = tmp_path / "agent.pptx"
    agent = PptAgent(None, NoopStore(), tmp_path)
    agent.render(request, _fallback_outline(request), output)
    normalized = agent.normalize_outline_for_layout(_fallback_outline(request))
    agent.validate(output, len(normalized["slides"]))

    presentation = Presentation(output)
    assert len(presentation.slides) > 8
    assert round(presentation.slide_width / presentation.slide_height, 2) == 1.78
    cover_text = "\n".join(
        shape.text for shape in presentation.slides[0].shapes if shape.has_text_frame
    )
    assert "AI Agent" in cover_text
    assert "Designed by ITN Agents" in cover_text
    assert presentation.slides[1].notes_slide.notes_text_frame.text.strip()


def test_render_supports_standard_ratio_without_notes(tmp_path: Path):
    request = make_request(slideCount=6, aspectRatio="4:3", includeSpeakerNotes=False, style="minimal")
    output = tmp_path / "standard.pptx"
    agent = PptAgent(None, NoopStore(), tmp_path)
    agent.render(request, _fallback_outline(request), output)

    presentation = Presentation(output)
    assert len(presentation.slides) > 6
    assert round(presentation.slide_width / presentation.slide_height, 2) == 1.33


def test_request_rejects_invalid_slide_count():
    import pytest
    with pytest.raises(ValueError):
        make_request(slideCount=3)


def test_agent_chooses_large_scale_for_broad_requirement():
    request = make_request(
        slideCount=0,
        requirement="制作覆盖全部章节的完整课程，需要内容详尽、包含案例、练习与综合复习",
    )
    agent = PptAgent(SizePlanningLlm(), NoopStore(), Path("."))
    assert asyncio.run(agent.choose_slide_count(request)) == 55
    assert agent._fallback_slide_count(request) == 55


def test_interaction_tutor_plans_agenda_and_scale(monkeypatch):
    monkeypatch.setattr(presentation_agent_module, "EnhancedChatAgent", FakeEnhancedTutor)
    request = make_request(
        slideCount=0,
        requirement="系统讲解离散数学的逻辑、集合关系和图论方法，并安排例题与练习",
    )
    plan = asyncio.run(PptAgent(TutorAgendaLlm(), NoopStore(), Path(".")).plan_agenda_with_tutor(request))
    assert plan["slideCount"] == 30
    assert len(plan["chapters"]) == 5
    assert [chapter["title"] for chapter in plan["chapters"]] == [
        "命题与联结词", "范式", "集合运算", "等价关系", "图论方法",
    ]
    assert plan["chapters"][0]["objective"] == "建立形式化推理基础"
    assert plan["chapters"][0]["keyTopics"] == ["命题与联结词"]
    assert plan["toolsUsed"] == []


def test_agenda_rejects_quality_requirements_as_chapters():
    request = make_request(
        title="传统Agent与LLM驱动Agent的区别｜交互课件",
        requirement="围绕主题生成课件，内容准确、结构清晰，适合课堂讲解与自主学习。",
        slideCount=12,
    )
    plan = asyncio.run(
        PptAgent(RequirementEchoAgendaLlm(), NoopStore(), Path(".")).plan_agenda_with_tutor(request)
    )
    titles = [chapter["title"] for chapter in plan["chapters"]]
    assert all(title not in {"内容准确", "结构清晰", "自主学习"} for title in titles)
    assert any("概念边界" in title for title in titles)


def test_comparison_fallback_uses_subject_instead_of_quality_words():
    request = make_request(
        title="传统Agent与LLM驱动Agent的区别｜交互课件",
        requirement="围绕主题生成课件，内容准确、结构清晰，适合课堂讲解与自主学习。",
        slideCount=12,
    )
    chapters = PptAgent._chapters_from_requirement(request, 12)
    titles = [chapter["title"] for chapter in chapters]
    assert titles[0].startswith("传统Agent与LLM驱动Agent的区别")
    assert "决策机制与工作流程对比" in titles
    assert all(title not in {"内容准确", "结构清晰", "自主学习"} for title in titles)


def test_tutor_agenda_is_preserved_as_section_architecture(monkeypatch):
    monkeypatch.setattr(presentation_agent_module, "EnhancedChatAgent", FakeEnhancedTutor)
    request = make_request(slideCount=12)
    chapters = asyncio.run(
        PptAgent(TutorAgendaLlm(), NoopStore(), Path(".")).plan_agenda_with_tutor(request)
    )["chapters"]
    normalized = PptAgent.normalize_outline_for_layout(_fallback_outline(request, chapters))
    agenda_headings = [
        point["heading"]
        for slide in normalized["slides"] if slide.get("kind") == "agenda"
        for point in slide["points"]
    ]
    section_titles = [
        slide["title"] for slide in normalized["slides"] if slide.get("kind") == "section"
    ]
    assert agenda_headings == ["命题逻辑", "集合与关系", "图论方法"]
    assert section_titles == agenda_headings


def test_broad_three_chapter_plan_expands_with_slide_scale():
    chapters = [
        {"title": "线性结构", "objective": "掌握线性结构", "keyTopics": ["顺序表", "链表", "栈", "队列"]},
        {"title": "树结构", "objective": "掌握树结构", "keyTopics": ["二叉树", "遍历", "堆", "并查集"]},
        {"title": "图与算法", "objective": "掌握图算法", "keyTopics": ["图存储", "遍历", "最短路", "最小生成树"]},
    ]
    expanded = PptAgent._expand_chapters(chapters, 40)
    assert len(expanded) >= 6
    assert {topic for chapter in expanded for topic in chapter["keyTopics"]} >= {
        "顺序表", "链表", "二叉树", "遍历", "最短路",
    }


def test_fallback_chapters_preserve_topics_named_by_user():
    request = make_request(
        slideCount=0,
        requirement="面向考研学生系统讲解数据结构，完整覆盖线性表、栈和队列、串、树与二叉树、图、查找、排序、算法复杂度和历年真题方法，内容详细且适合系统复习",
    )
    chapters = PptAgent._chapters_from_requirement(request, 40)
    titles = [chapter["title"] for chapter in chapters]
    assert len(chapters) >= 6
    assert "线性表" in titles
    assert "栈" in titles
    assert "树" in titles
    assert "查找" in titles
    assert "排序" in titles


def test_all_style_presets_render(tmp_path: Path):
    styles = ["academic", "business", "technology", "minimal", "vibrant", "education",
              "forest", "ocean", "terracotta", "editorial", "luxury", "dark"]
    for style in styles:
        request = make_request(slideCount=6, style=style)
        output = tmp_path / f"{style}.pptx"
        agent = PptAgent(None, NoopStore(), tmp_path)
        agent.render(request, _fallback_outline(request), output)
        agent.validate(output, len(agent.normalize_outline_for_layout(_fallback_outline(request))["slides"]))


def test_quality_gate_rejects_sparse_three_bullet_outline():
    request = make_request(slideCount=8, detailLevel="detailed")
    sparse = {"slides": [
        {"title": f"第{i + 1}页", "kind": "content", "points": ["短句一", "短句二", "短句三"]}
        for i in range(8)
    ]}
    assert not PptAgent._outline_is_rich(sparse, request)
    assert not PptAgent._outline_is_rich(_fallback_outline(request), request)


def test_technical_fallback_contains_code_and_explanation(tmp_path: Path):
    request = make_request(
        requirement="面向初学者详解 C 语言程序结构、变量、分支、循环与函数，并包含可运行代码",
        title="C 语言程序设计",
        slideCount=10,
        detailLevel="detailed",
    )
    outline = _fallback_outline(request)
    code_slides = [slide for slide in outline["slides"] if slide.get("kind") == "code"]
    assert code_slides
    assert "int main" in code_slides[0]["code"]["content"]
    assert len(code_slides[0]["code"]["annotations"]) >= 4

    output = tmp_path / "detailed-code.pptx"
    agent = PptAgent(None, NoopStore(), tmp_path)
    agent.render(request, outline, output)
    text = "\n".join(shape.text for slide in Presentation(output).slides for shape in slide.shapes if shape.has_text_frame)
    assert "int main" in text
    assert "运行结果" in text


def test_layout_normalization_clamps_complex_page_capacity():
    outline = {"slides": [{
        "title": "**很长的标题**" * 20,
        "kind": "case/practice",
        "takeaway": "总结" * 80,
        "points": [{"heading": "标题" * 20, "detail": "正文" * 200}] * 6,
        "code": {
            "caption": "代码说明" * 30,
            "content": "\n".join(f"line {index}" for index in range(30)),
            "annotations": ["注释" * 60] * 7,
            "output": "输出" * 100,
        },
        "process": [{"label": "节点" * 20, "detail": "说明" * 100}] * 8,
        "table": {"headers": list("ABCDE"), "rows": [["单元格" * 50] * 5] * 8},
    }]}
    slide = PptAgent.normalize_outline_for_layout(outline)["slides"][0]
    assert len(slide["points"]) == 4
    assert len(slide["code"]["content"].splitlines()) == 18
    assert len(slide["code"]["annotations"]) == 4
    assert len(slide["process"]) == 5
    assert len(slide["table"]["headers"]) == 4
    assert len(slide["table"]["rows"]) == 5
    assert "`" not in slide["title"]
    assert "**" not in slide["title"]
    assert slide["kind"] == "case"


def test_process_object_from_model_is_normalized_and_renderable(tmp_path: Path):
    request = make_request(slideCount=6, backgroundStyle="blue-white-minimal")
    outline = _fallback_outline(request)
    process_slide = next(slide for slide in outline["slides"] if slide.get("kind") == "process")
    process_slide["process"] = {"steps": [
        {"name": f"步骤{i}", "detail": f"说明{i}"} for i in range(1, 7)
    ]}
    normalized = PptAgent.normalize_outline_for_layout(outline)
    normalized_process = next(slide for slide in normalized["slides"] if slide.get("kind") == "process")
    assert len(normalized_process["process"]) == 5
    assert normalized_process["process"][0]["label"] == "步骤1"
    output = tmp_path / "object-process.pptx"
    PptAgent(None, NoopStore(), tmp_path).render(request, normalized, output)
    assert output.is_file()


def test_all_eleven_background_themes_have_bound_palettes():
    assert len(THEME_PALETTES) == 11
    assert palette_for("blue-conference") == ("102B50", "EDF4FA", "3F8ECD")
    assert palette_for("unknown") == palette_for("none")
    assert COVER_TEXT_OVERRIDES["watercolor"] == ("16232B", "243844")


def test_agent_fallback_uses_domain_specific_python_example():
    request = make_request(
        requirement="面向学生讲解 AI Agent 的规划、工具调用、记忆和错误恢复机制",
        title="AI Agent 工程实践",
        slideCount=10,
        detailLevel="detailed",
    )
    outline = _fallback_outline(request)
    assert any("工具调用" in slide["title"] for slide in outline["slides"])
    code_slide = next(slide for slide in outline["slides"] if slide.get("kind") == "code")
    assert code_slide["code"]["language"] == "python"
    assert "run_agent" in code_slide["code"]["content"]


def test_computer_architecture_fallback_uses_domain_topics():
    request = make_request(
        requirement="面向学生讲解计算机组成原理、冯诺依曼体系、指令周期、Cache和中断",
        title="计算机组成原理",
        slideCount=10,
        detailLevel="detailed",
    )
    titles = [slide["title"] for slide in _fallback_outline(request)["slides"]]
    assert "冯诺依曼体系结构" in titles
    assert "取指—译码—执行周期" in titles


def test_every_rendered_deck_has_agenda_as_second_slide(tmp_path: Path):
    request = make_request(slideCount=8)
    output = tmp_path / "with-agenda.pptx"
    PptAgent(None, NoopStore(), tmp_path).render(request, _fallback_outline(request), output)
    presentation = Presentation(output)
    assert presentation.slides[1].shapes[0].text == "目录"
    agenda_text = "\n".join(
        shape.text for shape in presentation.slides[1].shapes if shape.has_text_frame
    )
    assert "03" in agenda_text or "02" in agenda_text
    assert "本页结论" not in agenda_text
    assert "01" in agenda_text


def test_three_item_minimum_agenda_is_spread_vertically(tmp_path: Path):
    request = make_request(slideCount=6, backgroundStyle="clean-minimal")
    output = tmp_path / "three-item-agenda.pptx"
    PptAgent(None, NoopStore(), tmp_path).render(request, _fallback_outline(request), output)
    agenda = Presentation(output).slides[1]
    row_tops = sorted({
        round(shape.top / 914400, 2)
        for shape in agenda.shapes
        if shape.shape_type == 1 and round(shape.width / 914400, 1) > 8
    })
    assert row_tops == [1.93, 3.27, 4.61]


def test_large_deck_keeps_up_to_ten_sections_on_one_agenda_slide():
    request = make_request(slideCount=35)
    normalized = PptAgent.normalize_outline_for_layout(_fallback_outline(request))
    assert PptAgent.normalize_outline_for_layout(normalized) == normalized
    agendas = [slide for slide in normalized["slides"] if slide.get("kind") == "agenda"]
    sections = [slide for slide in normalized["slides"] if slide.get("kind") == "section"]
    assert len(sections) >= 6
    assert len(agendas) == 1
    assert sum(len(slide["points"]) for slide in agendas) == len(sections)
    assert all(len(slide["points"]) <= 10 for slide in agendas)
    assert [slide["startNumber"] for slide in agendas] == [1]


def test_more_than_ten_sections_still_paginate_safely():
    request = make_request(slideCount=70)
    chapters = [
        {"title": f"章节{index}", "objective": f"目标{index}", "keyTopics": [f"主题{index}"]}
        for index in range(1, 12)
    ]
    normalized = PptAgent.normalize_outline_for_layout(_fallback_outline(request, chapters))
    agendas = [slide for slide in normalized["slides"] if slide.get("kind") == "agenda"]

    assert len(agendas) == 2
    assert [len(slide["points"]) for slide in agendas] == [10, 1]
    assert [slide["startNumber"] for slide in agendas] == [1, 11]


def test_each_agenda_chapter_has_a_section_title_slide():
    request = make_request(slideCount=10)
    normalized = PptAgent.normalize_outline_for_layout(_fallback_outline(request))
    slides = normalized["slides"]
    agenda = slides[1]
    sections = [slide for slide in slides if slide.get("kind") == "section"]
    assert len(sections) == len(agenda["points"])
    for section in sections:
        section_index = slides.index(section)
        assert section_index >= 2
        assert slides[section_index + 1]["kind"] != "section"
        assert section["title"] == section["sectionTopics"][0]


def test_background_selection_is_stable_and_role_aware():
    first = select_backgrounds("blue-conference", "task-123")
    second = select_backgrounds("blue-conference", "task-123")
    assert first == second
    assert all(first.values())
    assert len(set(first.values())) >= 3


def test_cover_always_has_natural_subtitle_and_visible_copy_is_clean(tmp_path: Path):
    request = make_request(
        slideCount=8,
        backgroundStyle="abstract-gradient",
    )
    outline = _fallback_outline(request)
    outline["subtitle"] = ""
    outline["slides"][2]["takeaway"] = "本页重点：先理解概念 · 再完成案例"
    normalized = PptAgent.normalize_outline_for_layout(outline)
    assert normalized["subtitle"]
    assert "PPT Agent" not in normalized["subtitle"]
    assert "·" not in str(normalized)
    assert not any(
        str(slide.get("takeaway", "")).startswith(("本页重点", "本页结论"))
        for slide in normalized["slides"]
    )

    output = tmp_path / "clean-copy.pptx"
    PptAgent(None, NoopStore(), tmp_path).render(request, normalized, output)
    cover_text = "\n".join(
        shape.text for shape in Presentation(output).slides[0].shapes if shape.has_text_frame
    )
    assert normalized["subtitle"] in cover_text
    assert "Designed by ITN Agents" in cover_text


def test_plain_string_points_get_specific_card_headings():
    heading, detail = PptAgent._point_parts("明确下一步行动")
    assert heading == "明确下一步行动"
    assert detail == "明确下一步行动"
    assert heading != "重点"


def test_convert_math_to_unicode_greek_operators():
    result = PptAgent._convert_math_to_unicode(r"希腊字母 \alpha \beta \gamma \pi")
    assert "α" in result and "β" in result and "γ" in result and "π" in result
    assert r"\alpha" not in result


def test_convert_math_to_unicode_integral():
    result = PptAgent._convert_math_to_unicode(r"\int_a^b f(x) dx")
    assert "∫" in result
    assert r"\int" not in result
    assert "ₐ" in result or "_ₐ" in result


def test_convert_math_to_unicode_frac():
    result = PptAgent._convert_math_to_unicode(r"\frac{a}{b}")
    assert "/" in result
    assert r"\frac" not in result


def test_convert_math_to_unicode_sqrt():
    result = PptAgent._convert_math_to_unicode(r"\sqrt{x^2 + y^2}")
    assert "√" in result
    assert r"\sqrt" not in result


def test_convert_math_to_unicode_superscript():
    result = PptAgent._convert_math_to_unicode(r"x^2 + x^{ab}")
    assert "²" in result
    assert "^(ab)" in result or "x^(ab)" in result


def test_convert_math_to_unicode_subscript():
    result = PptAgent._convert_math_to_unicode(r"x_1 + x_{ij}")
    assert "₁" in result
    assert "ᵢ" in result or "_ᵢ" in result


def test_convert_math_to_unicode_relations():
    result = PptAgent._convert_math_to_unicode(r"\leq \geq \neq \approx")
    assert "≤" in result and "≥" in result and "≠" in result


def test_convert_math_to_unicode_delimiters():
    result = PptAgent._convert_math_to_unicode(r"内联公式 $E = mc^2$")
    assert "²" in result
    assert "$" not in result
    result2 = PptAgent._convert_math_to_unicode(r"$$ \int_0^1 x^2 dx $$")
    assert "∫" in result2
    assert "²" in result2


def test_convert_math_to_unicode_paren_delimiters():
    result = PptAgent._convert_math_to_unicode(r"\( \alpha + \beta \)")
    assert "α" in result and "β" in result
    result2 = PptAgent._convert_math_to_unicode(r"\[ \sum_{i=1}^n i \]")
    assert "∑" in result2


def test_normalize_outline_does_not_truncate_detail():
    outline = {"title": "测试", "subtitle": "测试", "slides": [
        {"kind": "concept", "title": "slide 1", "points": [{"heading": "A", "detail": "x" * 300}]},
        {"kind": "concept", "title": "slide 2", "takeaway": "y" * 200, "points": [{"heading": "B", "detail": "z" * 300}]},
    ]}
    normalized = PptAgent.normalize_outline_for_layout(outline)
    assert len(normalized["slides"][0]["points"][0]["detail"]) == 300
    assert normalized["slides"][1]["takeaway"] == "y" * 200


def test_normalize_outline_does_not_add_ellipsis():
    outline = {"title": "测试", "subtitle": "测试", "slides": [
        {"kind": "concept", "title": "t" * 100, "points": [{"heading": "h" * 100, "detail": "d"}]},
    ]}
    normalized = PptAgent.normalize_outline_for_layout(outline)
    assert "…" not in normalized["slides"][0]["points"][0]["heading"]
    assert "…" not in normalized["slides"][0]["title"]


def test_normalize_outline_turns_presenter_instructions_into_connected_speech():
    outline = {
        "title": "数据结构",
        "subtitle": "从概念到实践",
        "slides": [
            {"title": "数据结构", "kind": "title", "points": [], "speakerNotes": "介绍主题。"},
            {
                "title": "目录",
                "kind": "agenda",
                "startNumber": 3,
                "points": [
                    {"heading": "线性结构", "detail": "栈与队列"},
                    {"heading": "非线性结构", "detail": "树与图"},
                ],
                "speakerNotes": "介绍整份演示的章节结构。",
            },
            {
                "title": "线性结构",
                "kind": "section",
                "sectionNumber": 1,
                "sectionTopics": ["栈", "队列"],
                "speakerNotes": "进入第一章，说明学习目标。",
            },
            {
                "title": "栈的工作方式",
                "kind": "concept",
                "points": [
                    {"heading": "后进先出", "detail": "最后压入的元素最先弹出。"},
                    {"heading": "典型操作", "detail": "入栈和出栈都在栈顶完成。"},
                ],
                "speakerNotes": "展开讲解栈的核心机制。",
            },
            {
                "title": "队列的工作方式",
                "kind": "concept",
                "points": [
                    {"heading": "先进先出", "detail": "最早进入队列的元素最先离开。"},
                    {"heading": "实际应用", "detail": "任务调度会按照到达顺序处理请求。"},
                ],
                "speakerNotes": "本页需要引导学生理解队列，并建议讲解时配合示意图，让学生讨论具体案例。",
            },
            {
                "title": "总结与练习",
                "kind": "conclusion",
                "points": [{"heading": "练习", "detail": "判断操作序列是否合法。"}],
                "speakerNotes": "最后，我们一起回顾今天的核心内容，并完成一道练习。",
            },
        ],
    }

    slides = outline["slides"]
    PptAgent._ensure_spoken_speaker_notes(slides)

    assert slides[0]["speakerNotes"].startswith("大家好，今天我们围绕")
    assert "第3部分是线性结构" in slides[1]["speakerNotes"]
    assert "第4部分是非线性结构" in slides[1]["speakerNotes"]
    assert slides[2]["speakerNotes"] == "第1章聚焦线性结构，重点理解栈、队列。"
    assert not slides[3]["speakerNotes"].startswith(("首先来看", "接下来是"))
    assert "最后压入的元素最先弹出" in slides[3]["speakerNotes"]
    assert "最早进入队列的元素最先离开" in slides[4]["speakerNotes"]
    assert "引导学生" not in slides[4]["speakerNotes"]
    assert not slides[4]["speakerNotes"].startswith(("首先", "接下来", "随后", "然后"))
    assert "继续往下推进" not in "".join(slide["speakerNotes"] for slide in slides)
    for slide in slides:
        assert not re.match(r"^(介绍|讲解|说明|展开|强调|引导|进入)", slide["speakerNotes"])


def test_speaker_notes_keep_content_but_remove_long_page_transition():
    slides = [
        {"title": "章节", "kind": "section", "sectionNumber": 1, "sectionTopics": ["概念"]},
        {
            "title": "核心概念",
            "kind": "concept",
            "points": [{"heading": "定义", "detail": "这是页面可见内容。"}],
            "speakerNotes": (
                "现在我们讨论核心概念。这个概念决定了后续方法的适用范围，也能帮助我们判断常见误区。"
                "下一页将介绍具体案例。"
            ),
        },
        {
            "title": "总结",
            "kind": "conclusion",
            "points": [{"heading": "结论", "detail": "完成回顾。"}],
            "speakerNotes": "核心内容已经讲完，我们可以用这三个结论检查是否真正掌握，并据此安排后续练习。",
        },
    ]

    PptAgent._ensure_spoken_speaker_notes(slides)

    assert slides[1]["speakerNotes"].startswith("这个概念决定了")
    assert "这个概念决定了后续方法的适用范围" in slides[1]["speakerNotes"]
    assert "下一页" not in slides[1]["speakerNotes"]
    assert slides[2]["speakerNotes"].startswith("核心内容已经讲完")


def test_speaker_notes_do_not_force_sequence_word_on_each_body_page():
    slides = [
        {"title": "目录", "kind": "agenda", "points": [{"heading": "基础", "detail": "建立概念"}]},
        {
            "title": "符号主义",
            "kind": "concept",
            "points": [{"heading": "规则表示", "detail": "知识由符号和规则显式表达。"}],
            "speakerNotes": "首先，我们来看“符号主义”。它通过规则进行推理。",
        },
        {
            "title": "连接主义",
            "kind": "comparison",
            "points": [{"heading": "参数学习", "detail": "模型从数据中学习分布式表示。"}],
            "speakerNotes": "接下来，我们分析“连接主义”。它通过参数学习形成表示。",
        },
        {
            "title": "范式差异",
            "kind": "comparison",
            "points": [{"heading": "核心差异", "detail": "两种范式的知识表示方式不同。"}],
            "speakerNotes": "然后，我们讨论“范式差异”。差异集中在表示、学习和推理方式。",
        },
    ]

    PptAgent._ensure_spoken_speaker_notes(slides)

    assert slides[1]["speakerNotes"].startswith("它通过规则")
    assert slides[2]["speakerNotes"].startswith("它通过参数")
    assert slides[3]["speakerNotes"].startswith("差异集中")
    assert sum(
        bool(re.match(r"^(首先|接下来|随后|然后)", slide["speakerNotes"]))
        for slide in slides[1:]
    ) == 0


def test_transition_topic_fallback_uses_subject_story_not_request_instructions():
    request = PresentationRequest(
        requirement=(
            "面向职场新人制作《从符号主义到连接主义的范式转变》交互课件PPT，"
            "目标是RAG应用工程师。内容要覆盖核心概念、原理、方法、典型应用、练习检查与回顾，"
            "每一页有具体讲解，不用空泛标题；结合学习画像安排深度。"
        ),
        title="从符号主义到连接主义的范式转变 · 交互课件",
        slideCount=30,
        includeSpeakerNotes=True,
    )

    chapters = PptAgent._chapters_from_requirement(request, 30)
    titles = [chapter["title"] for chapter in chapters]

    assert any("符号主义如何表示与推理" in title for title in titles)
    assert any("连接主义如何学习表示" in title for title in titles)
    assert any("关键差异" in title for title in titles)
    assert all("目标是RAG应用工程师" not in title for title in titles)
    assert all("不用空泛标题" not in title for title in titles)


def test_section_and_first_content_slide_do_not_repeat_the_same_introduction():
    slides = [
        {
            "title": "背景与目标",
            "kind": "section",
            "sectionNumber": 1,
            "sectionTopics": ["背景与目标"],
        },
        {
            "title": "考研数学一：背景与学习目标",
            "kind": "concept",
            "points": [
                {
                    "heading": "考试定位",
                    "detail": "数学一主要面向对数学要求较高的工科专业。",
                },
            ],
            "speakerNotes": (
                "首先，我们来明确考研数学一的背景与学习目标。"
                "数学一主要面向对数学要求较高的工科专业，满分是一百五十分。"
            ),
        },
    ]

    PptAgent._ensure_spoken_speaker_notes(slides)

    assert slides[0]["speakerNotes"] == "第1章聚焦背景与目标。"
    assert slides[1]["speakerNotes"].startswith("数学一主要面向")
    assert "首先，我们来明确" not in slides[1]["speakerNotes"]


def test_math_one_fallback_agenda_uses_real_subject_chapters():
    request = make_request(
        requirement="考研数学一，全章节内容",
        title="考研数学一",
        slideCount=30,
        detailLevel="detailed",
    )

    chapters = PptAgent._chapters_from_requirement(request, 30)
    titles = [chapter["title"] for chapter in chapters]

    assert "函数、极限与连续" in titles
    assert "线性代数" in titles
    assert "概率论与数理统计" in titles
    assert "背景与目标" not in titles


def test_estimate_text_height_increases_with_length():
    h1 = PptAgent._estimate_text_height("短文本", 5.0, 10)
    h2 = PptAgent._estimate_text_height("短文本" * 50, 5.0, 10)
    assert h2 > h1


def test_estimate_text_height_narrow_box_more_lines():
    h_wide = PptAgent._estimate_text_height("abc" * 50, 6.0, 10)
    h_narrow = PptAgent._estimate_text_height("abc" * 50, 2.0, 10)
    assert h_narrow > h_wide

def test_goal_driven_request_is_json_serializable():
    request = PresentationRequest(
        requirement="围绕智能体工程生成一份目标驱动的个性化课件",
        learningGoalId="50049edd-566b-4ad7-aa33-7490756a4800",
        generationContext={"goal": {"title": "智能体工程"}},
    )
    payload = request.model_dump(mode="json", by_alias=True)
    assert payload["learningGoalId"] == "50049edd-566b-4ad7-aa33-7490756a4800"
    assert payload["generationContext"]["goal"]["title"] == "智能体工程"
