"""Meaningful regression: preview uses the saved PPTX canvas and excludes notes."""
import importlib.util
from pathlib import Path
import tempfile
import unittest

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
SPEC = importlib.util.spec_from_file_location('pptx_preview', ROOT / 'services/learning/renderer/ppt-renderer/pptx_preview.py')
MODULE = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(MODULE)


class PptxPreviewTests(unittest.TestCase):
    def test_saved_canvas_text_tables_and_geometry_without_notes(self):
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(3), Inches(1))
        box.text = '正文 <安全> & 原样展示'
        box.text_frame.paragraphs[0].font.size = Pt(18)
        box.text_frame.paragraphs[0].font.color.rgb = RGBColor.from_string('7047EB')
        table = slide.shapes.add_table(2, 2, Inches(1), Inches(3), Inches(5), Inches(2)).table
        table.cell(0, 0).text = '比较维度'
        table.cell(1, 1).text = '实际单元格'
        slide.notes_slide.notes_text_frame.text = 'SECRET_SPEAKER_NOTES_MUST_NOT_APPEAR'
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'example.pptx'
            presentation.save(path)
            preview = MODULE.preview_from_pptx(path)
        self.assertIn('正文 &lt;安全&gt; &amp; 原样展示', preview)
        self.assertIn('比较维度', preview)
        self.assertIn('实际单元格', preview)
        self.assertIn('x="96.000" y="192.000"', preview)
        self.assertIn('#7047EB', preview)
        self.assertNotIn('SECRET_SPEAKER_NOTES', preview)
        self.assertEqual(preview.count('data-slide='), 1)


if __name__ == '__main__':
    unittest.main()
