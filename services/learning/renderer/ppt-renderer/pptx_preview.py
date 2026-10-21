"""Portable preview of the generated PPTX canvas, using its actual shape geometry.

Never reads notes_slide: speaker notes belong to the PPTX notes part, not slides.
The template renderer uses text, solid shapes, pictures and tables, all covered here.
"""
from __future__ import annotations

import base64
import html
import io
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.enum.dml import MSO_COLOR_TYPE, MSO_FILL_TYPE
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

EMU_PER_PIXEL = 9525


def px(value):
    return float(value) / EMU_PER_PIXEL


def color(value, default='#26332d'):
    try:
        if value.type == MSO_COLOR_TYPE.RGB:
            return '#' + str(value.rgb)
    except (AttributeError, TypeError, ValueError):
        pass
    return default


def fill(value, default='none'):
    try:
        if value.type == MSO_FILL_TYPE.SOLID:
            return color(value.fore_color, default)
    except (AttributeError, TypeError, ValueError):
        pass
    return default


def font_style(font, fallback=None):
    fallback = fallback or font
    size = font.size or fallback.size
    name = font.name or fallback.name or 'Microsoft YaHei'
    bold = font.bold if font.bold is not None else fallback.bold
    return (f'font-family:{html.escape(name, quote=True)},sans-serif;'
            f'font-size:{float(size.pt) * 4 / 3 if size else 18.67:.3f}px;'
            f'font-weight:{700 if bold else 400};'
            f'font-style:{"italic" if font.italic else "normal"};'
            f'color:{color(font.color, color(fallback.color))}')


def text_frame(frame, x, y, width, height):
    paragraphs = []
    for paragraph in frame.paragraphs:
        align = {PP_ALIGN.CENTER: 'center', PP_ALIGN.RIGHT: 'right',
                 PP_ALIGN.JUSTIFY: 'justify'}.get(paragraph.alignment, 'left')
        runs = []
        for run in paragraph.runs:
            value = html.escape(run.text).replace('\n', '<br>').replace('\v', '<br>')
            runs.append(f'<span style="{font_style(run.font, paragraph.font)}">{value}</span>')
        if not runs:
            runs.append(html.escape(paragraph.text).replace('\v', '<br>') or '<br>')
        before = px(paragraph.space_before) if paragraph.space_before else 0
        after = px(paragraph.space_after) if paragraph.space_after else 0
        paragraphs.append(f'<p style="margin:{before:.2f}px 0 {after:.2f}px;'
                          f'text-align:{align};{font_style(paragraph.font)}">'
                          + ''.join(runs) + '</p>')
    anchor = {MSO_ANCHOR.MIDDLE: 'center', MSO_ANCHOR.BOTTOM: 'flex-end'}.get(frame.vertical_anchor, 'flex-start')
    padding = ' '.join(f'{px(v):.3f}px' for v in
                       (frame.margin_top, frame.margin_right, frame.margin_bottom, frame.margin_left))
    wrap = 'pre-wrap' if frame.word_wrap is not False else 'pre'
    return (f'<foreignObject x="{x:.3f}" y="{y:.3f}" width="{width:.3f}" height="{height:.3f}">'
            f'<div xmlns="http://www.w3.org/1999/xhtml" class="text-frame" '
            f'style="padding:{padding};justify-content:{anchor};white-space:{wrap}">'
            '<div class="text-content">' + ''.join(paragraphs) + '</div></div></foreignObject>')


def picture_markup(shape, cache):
    key = (shape.image.sha1, shape.crop_left, shape.crop_top, shape.crop_right, shape.crop_bottom)
    if key not in cache:
        blob = shape.image.blob
        mime = shape.image.content_type
        if any(key[1:]):
            with Image.open(io.BytesIO(blob)) as image:
                width, height = image.size
                image = image.crop((round(width * shape.crop_left), round(height * shape.crop_top),
                                    round(width * (1 - shape.crop_right)), round(height * (1 - shape.crop_bottom))))
                output = io.BytesIO()
                image.save(output, format='PNG')
                blob, mime = output.getvalue(), 'image/png'
        cache[key] = f'data:{mime};base64,' + base64.b64encode(blob).decode('ascii')
    return (f'<image x="{px(shape.left):.3f}" y="{px(shape.top):.3f}" '
            f'width="{px(shape.width):.3f}" height="{px(shape.height):.3f}" '
            f'preserveAspectRatio="none" href="{cache[key]}"/>')


def table_markup(shape):
    result = []
    y = px(shape.top)
    for row in shape.table.rows:
        x = px(shape.left)
        for index, cell in enumerate(row.cells):
            width, height = px(shape.table.columns[index].width), px(row.height)
            result.append(f'<rect x="{x:.3f}" y="{y:.3f}" width="{width:.3f}" height="{height:.3f}" '
                          f'fill="{fill(cell.fill, "#ffffff")}" stroke="#dbe2e8" stroke-width=".7"/>')
            result.append(text_frame(cell.text_frame, x, y, width, height))
            x += width
        y += px(row.height)
    return ''.join(result)


def shape_markup(shape, cache):
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        return picture_markup(shape, cache)
    if shape.has_table:
        return table_markup(shape)
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        raise ValueError('The presentation template emitted an unsupported grouped shape')
    x, y, width, height = map(px, (shape.left, shape.top, shape.width, shape.height))
    background = fill(shape.fill) if hasattr(shape, 'fill') else 'none'
    border = color(shape.line.color, 'none') if hasattr(shape, 'line') else 'none'
    stroke = px(shape.line.width) if hasattr(shape, 'line') and shape.line.width else .7
    geometry = shape._element.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}prstGeom')
    kind = geometry.get('prst') if geometry is not None else 'rect'
    opacity = ''
    alphas = shape._element.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/main}solidFill/{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr/{http://schemas.openxmlformats.org/drawingml/2006/main}alpha')
    if alphas:
        opacity = f' fill-opacity="{int(alphas[0].get("val")) / 100000:.3f}"'
    if kind == 'ellipse':
        backdrop = f'<ellipse cx="{x+width/2:.3f}" cy="{y+height/2:.3f}" rx="{width/2:.3f}" ry="{height/2:.3f}"'
    elif kind == 'chevron':
        points = f'{x},{y} {x+width*.65},{y} {x+width},{y+height/2} {x+width*.65},{y+height} {x},{y+height} {x+width*.35},{y+height/2}'
        backdrop = f'<polygon points="{points}"'
    else:
        radius = min(width, height) * .12 if 'round' in kind.lower() else 0
        backdrop = f'<rect x="{x:.3f}" y="{y:.3f}" width="{width:.3f}" height="{height:.3f}" rx="{radius:.3f}"'
    result = f'{backdrop} fill="{background}" stroke="{border}" stroke-width="{stroke:.3f}"{opacity}/>'
    if shape.has_text_frame:
        result += text_frame(shape.text_frame, x, y, width, height)
    return result


def preview_from_pptx(path: Path, title='演示文稿'):
    presentation = Presentation(path)
    width, height = px(presentation.slide_width), px(presentation.slide_height)
    slides, image_cache = [], {}
    for index, slide in enumerate(presentation.slides, 1):
        background = fill(slide.background.fill, '#ffffff')
        content = ''.join(shape_markup(shape, image_cache) for shape in slide.shapes)
        slides.append(f'<section class="slide" aria-label="第 {index} 页" data-slide="{index}">'
                      f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {width:.3f} {height:.3f}" '
                      f'role="img" aria-label="第 {index} 页"><rect width="100%" height="100%" '
                      f'fill="{background}"/>{content}</svg></section>')
    css = ('body{margin:0;padding:16px;background:#ececf2;font-family:"Microsoft YaHei",sans-serif}'
           '.slide{max-width:1280px;margin:0 auto 18px;background:white;box-shadow:0 4px 18px #16112a14}'
           'svg{display:block;width:100%;height:auto}.text-frame{display:flex;flex-direction:column;'
           'box-sizing:border-box;width:100%;height:100%;overflow:hidden;line-height:1.15}'
           '.text-content{width:100%;overflow-wrap:break-word}p{padding:0;line-height:1.15}'
           '@media(max-width:640px){body{padding:6px}.slide{margin-bottom:8px}}')
    return ('<!doctype html><html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" '
            'content="width=device-width,initial-scale=1"><title>' + html.escape(title) + '</title><style>'
            + css + '</style></head><body>' + ''.join(slides) + '</body></html>')
