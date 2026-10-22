"""Selective multimodal layout review for high-risk presentation slides."""

from __future__ import annotations

import asyncio
import copy
import os
import shutil
import subprocess
import tempfile
import threading
from pathlib import Path
from typing import Any

from pptx import Presentation
from app.config import Settings
from app.llm import LLMClient

# PowerPoint COM on Windows cannot open the same file concurrently.
# Serialize COM rendering so parallel review calls do not conflict.
_pptx_com_lock = threading.Lock()


COMPLEX_KINDS = {"code", "table", "process", "timeline", "comparison", "case", "practice", "conclusion"}
ALLOWED_ISSUES = {"overlap", "overflow", "cutoff", "crowded", "misalignment", "low_contrast"}
ALLOWED_COVER_PLACEMENTS = {"left-top", "left-middle", "center", "right-top", "right-middle"}


class PresentationVisualReviewer:
    def __init__(self, llm: LLMClient, settings: Settings):
        self.llm = llm
        self.settings = settings

    def select_candidates(self, outline: dict[str, Any]) -> list[int]:
        scored: list[tuple[int, int]] = []
        slides = outline.get("slides") if isinstance(outline, dict) else []
        for index, slide in enumerate(slides or []):
            if not isinstance(slide, dict):
                continue
            kind = str(slide.get("kind", "content"))
            if kind not in COMPLEX_KINDS:
                continue
            score = 2
            points = slide.get("points") or []
            visible_chars = len(str(points))
            if visible_chars > 420:
                score += 3
            elif visible_chars > 260:
                score += 1
            if kind == "code":
                code = slide.get("code") or {}
                score += max(0, str(code.get("content", "")).count("\n") - 10) // 3
                score += max(0, len(code.get("annotations") or []) - 3)
            elif kind == "table":
                table = slide.get("table") or {}
                score += (len(table.get("headers") or []) * len(table.get("rows") or [])) // 8
            elif kind in {"process", "timeline"}:
                score += max(0, len(slide.get("process") or []) - 3)
            scored.append((score, index))
        scored.sort(key=lambda item: (-item[0], item[1]))
        return [index for _, index in scored[:self.settings.presentation_visual_review_max_slides]]

    async def review(self, path: Path, outline: dict[str, Any]) -> list[dict[str, Any]]:
        if not self.settings.presentation_visual_review_enabled or not self.settings.mimo_api_key:
            return []
        candidates = self.select_candidates(outline)
        if not candidates:
            return []
        workspace = Path(tempfile.mkdtemp(prefix="itn-ppt-review-"))
        try:
            images = await asyncio.to_thread(self._render_slides, path, candidates, workspace)
            if not images:
                return []
            result = await self.llm.complete_vision_json(
                "你是严格的PPT排版质检员，只判断可见的排版缺陷，不评价内容质量。不要把正常留白误判为问题。",
                self._review_prompt(candidates),
                [str(image) for image in images],
                {"slides": []},
                model=self.settings.presentation_visual_review_model,
                provider="mimo-legacy",
            )
            issues = self._normalize_issues(result, candidates)
            return issues
        finally:
            shutil.rmtree(workspace, ignore_errors=True)

    async def review_cover(self, path: Path, title: str, subtitle: str) -> dict[str, Any]:
        """Use one low-detail vision call to tune only cover typography."""
        if not self.settings.presentation_visual_review_enabled or not self.settings.mimo_api_key:
            return {}
        workspace = Path(tempfile.mkdtemp(prefix="itn-ppt-cover-review-"))
        try:
            images = await asyncio.to_thread(self._render_slides, path, [0], workspace)
            if not images:
                return {}
            result = await self.llm.complete_vision_json(
                "你是PPT封面视觉设计师。只根据封面截图和背景构图判断文字的对比色与最佳位置，不改写标题或副标题。",
                (
                    f"封面标题：{title}\n封面副标题：{subtitle}\n"
                    "必须重新判断文字应左对齐、居中还是右对齐：如果画面主体或装饰在两侧、中央留白完整，"
                    "优先选择center；如果只有左侧或右侧存在安全留白，则选择相应位置。"
                    "不要根据主题名称套固定规则。位置不得遮挡背景主体，标题必须醒目，副标题必须可读。"
                    "严格返回JSON：{\"placement\":\"left-top|left-middle|center|right-top|right-middle\","
                    "\"titleColor\":\"6位十六进制颜色\",\"subtitleColor\":\"6位十六进制颜色\","
                    "\"titleSize\":40到48的整数}。不要返回其他字段。"
                ),
                [str(images[0])],
                {},
                model=self.settings.presentation_visual_review_model,
                provider="mimo-legacy",
            )
            return self._normalize_cover_design(result)
        finally:
            shutil.rmtree(workspace, ignore_errors=True)

    @staticmethod
    def _normalize_cover_design(result: dict[str, Any]) -> dict[str, Any]:
        if not isinstance(result, dict):
            return {}
        placement = str(result.get("placement", ""))
        if placement not in ALLOWED_COVER_PLACEMENTS:
            return {}

        def color(value: Any, fallback: str) -> str:
            text = str(value or "").strip().lstrip("#").upper()
            return text if len(text) == 6 and all(char in "0123456789ABCDEF" for char in text) else fallback

        try:
            title_size = max(40, min(48, int(result.get("titleSize", 44))))
        except (TypeError, ValueError):
            title_size = 44
        return {
            "placement": placement,
            "titleColor": color(result.get("titleColor"), "FFFFFF"),
            "subtitleColor": color(result.get("subtitleColor"), "E5E7EB"),
            "titleSize": title_size,
        }

    @staticmethod
    def _review_prompt(candidates: list[int]) -> str:
        page_numbers = [index + 1 for index in candidates]
        return (
            f"图片依次对应PPT页码 {page_numbers}。只报告确定存在的问题：模块互相遮挡、文字溢出卡片、"
            "文字被裁切、内容过度拥挤、明显错位、低对比度。忽略普通留白、风格偏好和内容建议。"
            "编号圆点压在卡片边缘、页码、装饰形状覆盖卡片边框属于正常设计，不要判为重叠。"
            "严格返回JSON：{\"slides\":[{\"slide\":页码,\"severity\":\"medium|high\","
            "\"confidence\":0到1的小数,\"issues\":[\"overlap|overflow|cutoff|crowded|misalignment|low_contrast\"]}]}。"
            "没有明确问题返回 {\"slides\":[]}。"
        )

    @staticmethod
    def _normalize_issues(result: dict[str, Any], candidates: list[int]) -> list[dict[str, Any]]:
        allowed_pages = {index + 1 for index in candidates}
        normalized = []
        for item in result.get("slides", []) if isinstance(result, dict) else []:
            if not isinstance(item, dict):
                continue
            try:
                slide = int(item.get("slide"))
            except (TypeError, ValueError):
                continue
            issues = [value for value in item.get("issues", []) if value in ALLOWED_ISSUES]
            try:
                confidence = float(item.get("confidence", 0))
            except (TypeError, ValueError):
                confidence = 0
            if (
                slide not in allowed_pages
                or not issues
                or item.get("severity") != "high"
                or confidence < 0.85
            ):
                continue
            normalized.append({
                "slide": slide,
                "severity": "high",
                "confidence": confidence,
                "issues": issues,
            })
        return normalized

    @staticmethod
    def _render_slides(path: Path, candidates: list[int], output_dir: Path) -> list[Path]:
        if os.name != "nt":
            return PresentationVisualReviewer._render_slides_libreoffice(path, candidates, output_dir)
        try:
            import pythoncom
            import win32com.client
        except ImportError:
            return []
        images: list[Path] = []
        output_dir.mkdir(parents=True, exist_ok=True)
        with _pptx_com_lock:
            pythoncom.CoInitialize()
            app = None
            deck = None
            try:
                app = win32com.client.DispatchEx("PowerPoint.Application")
                deck = app.Presentations.Open(str(path.resolve()), True, False, False)
                for index in candidates:
                    image = (output_dir / f"slide-{index + 1:02d}.png").resolve()
                    # Export above thumbnail resolution so small Chinese glyphs do not rasterize
                    # into dark blocks that the vision model could mistake for overlap/cutoff.
                    deck.Slides(index + 1).Export(str(image), "PNG", 1440, 810)
                    if image.is_file():
                        images.append(image)
            except Exception:
                return []
            finally:
                if deck is not None:
                    deck.Close()
                if app is not None:
                    app.Quit()
                pythoncom.CoUninitialize()
        return images

    @staticmethod
    def _render_slides_libreoffice(path: Path, candidates: list[int], output_dir: Path) -> list[Path]:
        soffice = shutil.which("soffice")
        pdftoppm = shutil.which("pdftoppm")
        if not soffice or not pdftoppm:
            return []
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir", str(output_dir), str(path.resolve())],
                check=True, capture_output=True, timeout=90,
            )
            pdf = output_dir / f"{path.stem}.pdf"
            if not pdf.is_file():
                return []
            images = []
            for index in candidates:
                prefix = output_dir / f"slide-{index + 1:02d}"
                subprocess.run(
                    [pdftoppm, "-f", str(index + 1), "-l", str(index + 1), "-singlefile",
                     "-png", "-r", "135", str(pdf), str(prefix)],
                    check=True, capture_output=True, timeout=30,
                )
                image = prefix.with_suffix(".png")
                if image.is_file():
                    images.append(image)
            return images
        except (OSError, subprocess.SubprocessError):
            return []

    @staticmethod
    def compact_outline(outline: dict[str, Any], issues: list[dict[str, Any]]) -> dict[str, Any]:
        compacted = copy.deepcopy(outline)
        for issue in issues:
            slide = compacted["slides"][issue["slide"] - 1]
            for point in slide.get("points", []):
                if isinstance(point, dict) and "detail" in point:
                    detail = str(point["detail"])
                    point["detail"] = detail[:500]
            if isinstance(slide.get("code"), dict):
                code = slide["code"]
                code["annotations"] = [str(value)[:200] for value in (code.get("annotations") or [])[:4]]
                lines = str(code.get("content", "")).splitlines()
                code["content"] = "\n".join(lines[:30])
                code["output"] = str(code.get("output", ""))[:500]
            if isinstance(slide.get("table"), dict):
                rows = slide["table"].get("rows") or []
                if isinstance(rows, dict):
                    rows = rows.get("rows") or rows.get("items") or []
                slide["table"]["rows"] = rows[:10] if isinstance(rows, list) else []
            process = slide.get("process")
            if isinstance(process, dict):
                process = process.get("steps") or process.get("items") or process.get("nodes") or []
            if isinstance(process, list):
                slide["process"] = process[:10]
                for node in slide["process"]:
                    if isinstance(node, dict) and "detail" in node:
                        node["detail"] = str(node["detail"])[:500]
        return compacted

    @staticmethod
    def validate_bounds(path: Path, slide_indexes: list[int]) -> None:
        presentation = Presentation(path)
        for index in slide_indexes:
            for shape in presentation.slides[index].shapes:
                if shape.left < 0 or shape.top < 0:
                    raise RuntimeError(f"视觉修正后第 {index + 1} 页存在越界元素")
                if shape.left + shape.width > presentation.slide_width:
                    raise RuntimeError(f"视觉修正后第 {index + 1} 页存在横向越界元素")
                if shape.top + shape.height > presentation.slide_height:
                    raise RuntimeError(f"视觉修正后第 {index + 1} 页存在纵向越界元素")
